#!/usr/bin/env python3
"""政令指定都市の区の団体コード表を作る（区ごとの地域ページの土台）。

照合の正典（krefuge の muni_vintage）には区のコードが無い（14101 等は0件）。
総務省「全国地方公共団体コード」の Excel にある「政令指定都市」シートから、
区だけを取り出して data/seirei_wards.json に保存する。

  cd /home/kojima/work/khazard && /usr/bin/python3 scripts/fetch_wards.py

出典: 総務省 全国地方公共団体コード（令和6年1月1日更新）
  https://www.soumu.go.jp/denshijiti/code.html
  （2026-10-10 取得。浜松市の再編後＝中央区・浜名区・天竜区を含む20市171区）
"""
import io, json, os, re, sys
import requests
import openpyxl

SRC_PAGE = "https://www.soumu.go.jp/denshijiti/code.html"
XLSX = "https://www.soumu.go.jp/main_content/000925835.xlsx"
OUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "data", "seirei_wards.json")


def main():
    r = requests.get(XLSX, timeout=60, headers={"User-Agent": "khazard/1.0 (kurage.exbridge.jp)"})
    r.raise_for_status()
    wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True)
    ws = next(w for w in wb.worksheets if "政令指定都市" in w.title)
    cities, wards = {}, {}
    for row in ws.iter_rows(values_only=True):
        if not row or row[0] is None or not re.fullmatch(r"\d{5,6}", str(row[0]).strip()):
            continue
        code = str(row[0]).strip().zfill(6)[:5]     # 6桁目は検査数字
        pref, name = str(row[1]).strip(), str(row[2] or "").strip()
        if name.endswith("市"):
            cities[code] = name
    for row in ws.iter_rows(values_only=True):
        if not row or row[0] is None or not re.fullmatch(r"\d{5,6}", str(row[0]).strip()):
            continue
        code = str(row[0]).strip().zfill(6)[:5]
        pref, name = str(row[1]).strip(), str(row[2] or "").strip()
        if not name.endswith("区"):
            continue
        city = next((c for c, n in cities.items() if c[:2] == code[:2] and name.startswith(n)), None)
        if not city:
            print("市を特定できない区:", code, name, file=sys.stderr)
            continue
        wards[code] = {"pref": pref, "city_code": city, "city": cities[city],
                       "ward": name[len(cities[city]):], "name": name}
    out = {"_source": "総務省 全国地方公共団体コード（%s）" % ws.title, "_url": SRC_PAGE,
           "_xlsx": XLSX, "cities": cities, "wards": wards}
    json.dump(out, open(OUT, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print(f"政令市 {len(cities)} / 区 {len(wards)} -> {OUT}")


if __name__ == "__main__":
    main()
